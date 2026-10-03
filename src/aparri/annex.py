"""Phase 10E: technical annex for the thesis appendix -> ``outputs/reports/Technical_Annex_GIS.docx`` (+ .md).

Collects what a panel member or the next student needs to repeat the work: inputs with checksums, software versions,
parameters, commands, results, tests and sensitivity, limitations.  Everything is read from the repository's own
outputs, so it can be regenerated at any time with ``make annex``.
"""
from __future__ import annotations

import json
import subprocess
import sys

import numpy as np
import pandas as pd

from .audit_report import md
from .io import REPO_ROOT, ensure_dir, load_config, repo_path


def latest_report(step: str) -> dict:
    files = sorted((REPO_ROOT / "outputs" / "reports").glob(f"run_{step}_*.json"))
    return json.loads(files[-1].read_text()) if files else {}


def n_tests() -> str:
    r = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q", "-o", "addopts="], capture_output=True, text=True, cwd=REPO_ROOT)
    ids = [l for l in r.stdout.splitlines() if "::" in l]
    return f"{len(ids)} automated tests"


def build_md(cfg: dict) -> str:
    T = lambda n: pd.read_csv(repo_path(f"outputs/tables/{n}"))
    th = T("Table_4_2_barangay_results.csv")
    ext = T("Table_4_2_extended.csv")
    runs = T("sensitivity_runs.csv")[["scenario", "description", "pct_Low", "pct_Medium", "pct_High"]]
    rob = T("sensitivity_robustness.csv")
    pri = T("priority_ranking_field_inspection.csv")
    versions = latest_report("transects").get("versions", {})
    manifest = [l.split("  ", 1) for l in (REPO_ROOT / "data/raw/MANIFEST.sha256").read_text().splitlines()]
    key = [m for m in manifest if any(k in m[1] for k in ("Shoreline Changes.shp", "Risk_1990_2025.shp", "Aparri_Brgy_Study.shp", "Aparri Barangays.shp", "Code.docx", "EPR_1990_2025.tif", "Shoreline_2025.tif", "Statistics_1990_2025.csv"))]
    keytab = pd.DataFrame({"file": [k[1] for k in key], "sha256 (first 16)": [k[0][:16] for k in key]})
    params = {
        "Working CRS": f"EPSG:{cfg['crs_work']} (WGS 84 / UTM 51N, metres)", "Years": cfg["years"], "Shoreline set (primary)": cfg["shoreline_set"],
        "Baseline": f"1990 open-coast shoreline, smoothed {cfg['baseline']['smooth_m']} m, offset landward by max excursion + {cfg['baseline']['offset_landward_m']} m",
        "Transects": f"every {cfg['transects']['spacing_m']} m, {cfg['transects']['length_seaward_m']} m seaward / {cfg['transects']['length_landward_m']} m landward, multi-hit rule: {cfg['transects']['multi_hit_rule']}",
        "Open-coast rule": f"within {cfg['scope']['hull_tolerance_m']} m of the concave-hull outline (ratio {cfg['scope']['concave_ratio']}); estuarine banks excluded",
        "Cleaning": f"snap {cfg['clean']['snap_tol_m']} m; loops < {cfg['clean']['min_loop_m']} m removed; duplicate stretches ({cfg['clean']['duplicate_cover']:.0%} within {cfg['clean']['duplicate_tol_m']} m) removed",
        "Class limits": f"Low < {cfg['risk']['low_max']}, Medium {cfg['risk']['low_max']}–{cfg['risk']['medium_max']} (inclusive), High > {cfg['risk']['medium_max']} m/yr; accretion -> {cfg['risk']['accretion_class']} (tagged)",
        "Hectare strip": f"{cfg['risk']['strip_width_m']} m", "Barangay join tolerance": f"{cfg['barangay_join_tolerance_m']} m",
        "Acquisition dates": "NOT PROVIDED (EPR over 35 calendar years)", "Positional uncertainty": "NOT PROVIDED (what-if ±30 m only for the confidence label)",
    }
    ptab = pd.DataFrame({"parameter": list(params), "value": [str(v) for v in params.values()]})
    vtab = pd.DataFrame({"package": list(versions), "version": list(versions.values())})
    return f"""# Technical Annex — GIS analysis of shoreline change, Aparri, Cagayan  [DRAFT – FOR STUDENT REVIEW]

*Generated {pd.Timestamp.now().strftime('%Y-%m-%d')} by `make annex` from the repository. Purpose: let a reader repeat and check the GIS work.*

## 1. How to repeat the analysis
```
make setup        # creates .venv and installs the dependencies from pyproject.toml
make test         # {n_tests()}
make all          # audit, audit-gee, clean-lines, raster-lines, transects, risk, tables, maps, sensitivity, web
make priority qgis drafts annex survey
```
Every step reads `config/config.yaml` (all parameters, with comments) and writes `outputs/reports/run_<step>_<time>.json` (parameters, input checksums, software versions, counts, warnings).

## 2. Input data (read-only, checksums)
{md(keytab)}

All {len(manifest)} files in `data/raw/` are listed with SHA-256 in `data/raw/MANIFEST.sha256`; the test suite fails if any raw file changes.

## 3. Software
{md(vtab)}

## 4. Parameters used
{md(ptab)}

## 5. Results
**Table 4.2** (kilometres of shoreline are in the extended table):

{md(th.round(2), floatfmt='{:,.2f}')}

{md(ext[['Barangay', 'Low risk (km)', 'Medium risk (km)', 'High risk (km)', 'Unclassified (km)', 'n transects', 'n valid', 'n flagged', '% eroding', '% classes that may flip']].round(2), floatfmt='{:,.2f}')}

Figures: `outputs/maps/Fig_4_1_v2_multitemporal_shorelines.png`, `Fig_4_2_v2_erosion_risk_map.png`, `Fig_4_3_epr_profile_along_coast.png`, `Fig_legacy_vs_rebuilt.png`.

Field-inspection priority (erosion information only; weights in config):

{md(pri.round(2), floatfmt='{:,.2f}')}

## 6. Quality assurance
* Automated tests: {n_tests()} — synthetic coasts of known change (straight 2.000 m/yr retreat, 1.500 m/yr advance, circular arcs, missing years, loops, mirrored orientation), cleaning rules, class boundaries (exactly 2.0 and 5.0 are Medium), survey statistics against hand computations, web-map smoke test.
* Sensitivity (share of valid transects, %):

{md(runs.round(1), floatfmt='{:,.1f}')}

* Stability of the barangay classes across runs:

{md(rob.round(0), floatfmt='{:,.0f}')}

## 7. Findings about the legacy outputs (summary; full text in `docs/01_data_audit.md`, `docs/01b_gee_audit.md`)
The supplied risk polygons, Earth Engine barangay statistics and raster change layers could not be reproduced or were not valid for the thesis purpose; they were not used in any result.

## 8. Limitations
See `docs/02_method.md` §7: erosion rate only; 30 m imagery with no recorded dates, tide or measured error; river mouth excluded; no field verification yet.

## 9. File index
`config/` parameters · `src/aparri/` code · `data/raw/` inputs (read-only) · `data/interim/` cleaned shorelines + review package · `data/processed/` transects and segments · `outputs/tables|maps|layers|web|field_validation|reports/` results · `docs/` audit, method, drafts, Q&A.
"""


def run() -> None:  # pragma: no cover
    cfg = load_config()
    text = build_md(cfg)
    out = ensure_dir("outputs/reports")
    (out / "Technical_Annex_GIS.md").write_text(text, encoding="utf-8")
    # DOCX rendering: simple, readable, thesis-fonts
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"; doc.styles["Normal"].font.size = Pt(10)
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("# "):
            doc.add_heading(l[2:], 0)
        elif l.startswith("## "):
            doc.add_heading(l[3:], 1)
        elif l.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            hdr = [c.strip() for c in l.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip("|").split("|")]); j += 1
            t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
            for k, h in enumerate(hdr):
                t.rows[0].cells[k].text = h
                for r in t.rows[0].cells[k].paragraphs[0].runs:
                    r.bold = True; r.font.size = Pt(8)
            for rw in rows:
                cells = t.add_row().cells
                for k, v in enumerate(rw[:len(hdr)]):
                    cells[k].text = v
                    for r in cells[k].paragraphs[0].runs:
                        r.font.size = Pt(8)
            i = j; continue
        elif l.startswith("```"):
            j = i + 1; blk = []
            while j < len(lines) and not lines[j].startswith("```"):
                blk.append(lines[j]); j += 1
            p = doc.add_paragraph("\n".join(blk)); 
            for r in p.runs:
                r.font.name = "Courier New"; r.font.size = Pt(8)
            i = j + 1; continue
        elif l.strip():
            doc.add_paragraph(l.replace("**", "").replace("`", ""))
        i += 1
    for fig in ("Fig_4_1_v2_multitemporal_shorelines.png", "Fig_4_2_v2_erosion_risk_map.png", "Fig_4_3_epr_profile_along_coast.png"):
        p = repo_path(f"outputs/maps/{fig}")
        if p.exists():
            doc.add_paragraph(fig).runs[0].bold = True
            doc.add_picture(str(p), width=Inches(6.3))
    doc.save(out / "Technical_Annex_GIS.docx")
    print("annex written")


if __name__ == "__main__":  # pragma: no cover
    run()
