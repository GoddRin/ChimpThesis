"""Phase 10F: a short slide deck (PPTX) for the GIS part of the defense -> ``outputs/defense/Aparri_GIS_defense_slides.pptx``.

Numbers are read from the pipeline tables (via ``drafts.facts``) so the slides cannot disagree with Table 4.2.
The deck is a starting point for the students; speaker notes carry the talking points and the honest caveats.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .drafts import facts
from .io import ensure_dir, load_config, repo_path

NAVY = RGBColor(0x1F, 0x3A, 0x5F); GREY = RGBColor(0x55, 0x55, 0x55); ORANGE = RGBColor(0xE0, 0x8A, 0x00); RED = RGBColor(0xB0, 0x20, 0x20)
W, H = Inches(13.333), Inches(7.5)


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=NAVY, align=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.text = line; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color; p.font.name = "Calibri"
        p.space_after = Pt(6)
        if align:
            p.alignment = align
    return tb


def title(slide, text, sub=None):
    add_text(slide, Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9), text, 30, True)
    if sub:
        add_text(slide, Inches(0.5), Inches(1.05), Inches(12.3), Inches(0.5), sub, 16, False, GREY)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def pic(slide, path, x, y, w=None, h=None):
    p = repo_path(path)
    if p.exists():
        slide.shapes.add_picture(str(p), x, y, w, h)


def table(slide, df: pd.DataFrame, x, y, w, h, font=11):
    shp = slide.shapes.add_table(len(df) + 1, len(df.columns), x, y, w, h)
    t = shp.table
    for j, c in enumerate(df.columns):
        t.cell(0, j).text = str(c)
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, v in enumerate(r.tolist()):
            t.cell(i, j).text = f"{v:.2f}" if isinstance(v, float) else str(v)
    for i in range(len(df) + 1):
        for j in range(len(df.columns)):
            for p in t.cell(i, j).text_frame.paragraphs:
                p.font.size = Pt(font); p.font.name = "Calibri"
                if j > 0:
                    p.alignment = PP_ALIGN.CENTER
    return t


def build(cfg: dict) -> Presentation:
    f = facts(cfg)
    ext = f["ext"]; u = f["u"]
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    add_text(s, Inches(0.7), Inches(2.0), Inches(12), Inches(1.6), "Shoreline change and erosion-rate classes along eight coastal barangays of Aparri, Cagayan, 1990–2025", 36, True)
    add_text(s, Inches(0.7), Inches(4.0), Inches(12), Inches(1), ["GIS component of the thesis “Development of Construction Management Framework for Shoreline Protection Structures …”", "Cagayan State University – Carig Campus · [DRAFT – FOR STUDENT REVIEW]"], 18, False, GREY)
    notes(s, "Introduce the GIS part only. Say at the start that the GIS analysis was rebuilt in a documented, scripted workflow and that the figures differ from the proposal-stage figures (see the adviser memo).")

    s = prs.slides.add_slide(blank); title(s, "What we measured and why")
    add_text(s, Inches(0.6), Inches(1.5), Inches(6.2), Inches(5), [
        "Question (Objective 1): which parts of the coast erode slowly, moderately or fast?",
        "Measure: End Point Rate (EPR) = shoreline movement 1990 → 2025 ÷ 35 years; negative = landward = erosion.",
        "Classes (Table 3.1): Low < 2 · Medium 2–5 · High > 5 m/yr.",
        f"Study area: open sea-facing coast of 8 barangays, {f['km_tot']['Classified shoreline (km)'] + f['km_tot']['Unclassified (km)']:.1f} km, from Bulala Sur to Paddaya.",
        "The classes describe the erosion RATE only — not exposure or vulnerability."], 18)
    pic(s, "outputs/maps/method_transect_diagram.png", Inches(7.0), Inches(1.8), w=Inches(6.0))
    notes(s, "Define EPR in one sentence with the diagram: two distances measured along the same transect, subtract, divide by 35 years.")

    s = prs.slides.add_slide(blank); title(s, "How the numbers were produced", "Reproducible workflow – every step logged; all parameters in one config file")
    add_text(s, Inches(0.6), Inches(1.7), Inches(12), Inches(5), [
        "1  Shoreline lines for 1990, 2000, 2010, 2020, 2025 → cleaned (duplicate stretch, loops, fragments) with a logged script",
        "2  Open coast separated from the Cagayan River banks by a fixed geometric rule",
        f"3  Baseline on land behind all years; {f['n_tr']} transects every {cfg['transects']['spacing_m']} m ({f['n_valid']} valid)",
        "4  NSM, EPR (1990–2025) and LRR (all years); classes from Table 3.1",
        "5  Barangay summary (Table 4.2): kilometres of shoreline per class (+ hectares of a 100 m strip)",
        "[STUDENTS: add one line on the imagery source, sensors and dates]"], 19)
    notes(s, "If asked about the source of the shorelines, answer from docs/04_defense_qa.md Q7 — fill in the origin truthfully.")

    s = prs.slides.add_slide(blank); title(s, "How we know the method is right", "Checks with known answers, a second data set, and sensitivity")
    add_text(s, Inches(0.6), Inches(1.7), Inches(6.2), Inches(5), [
        "Synthetic coasts with known change: 2.000 m/yr retreat → EPR −2.000; 1.500 m/yr advance → +1.500; curved coast within 0.01 m/yr",
        f"Independent shoreline set (from the Earth Engine masks): EPR r = {f['src_r']:.2f}; same class for {f['src_agree']:.0f} % of {f['src_n']} transects",
        f"Spacing, smoothing, baseline offset change any class share by ≤ {f['max_shift']:.1f} percentage points",
        f"Shifting the class limits by ±0.5 m/yr moves the Medium share from {f['runs'].loc['base','pct_Medium']:.0f} % to {f['runs'].loc['thr_+0.5','pct_Medium']:.0f}–{f['runs'].loc['thr_-0.5','pct_Medium']:.0f} %"], 17)
    pic(s, "outputs/maps/Fig_6_sensitivity_class_shares.png", Inches(7.0), Inches(2.0), w=Inches(6.0))
    notes(s, "The honest message: the technical choices do not matter; the class limits are conventions and the coast sits near the 2 m/yr limit.")

    s = prs.slides.add_slide(blank); title(s, "Figure 4.1 — shorelines 1990–2025")
    pic(s, "outputs/maps/Fig_4_1_v2_multitemporal_shorelines.png", Inches(2.35), Inches(1.2), h=Inches(6.1))
    notes(s, "Point to inset A: at Bulala Norte the 1990 line (purple) lies well seaward of the later lines. Inset B: the river mouth, where lines cross — not classified.")

    s = prs.slides.add_slide(blank); title(s, "Figure 4.2 — erosion-rate classes")
    pic(s, "outputs/maps/Fig_4_2_v2_erosion_risk_map.png", Inches(2.35), Inches(1.2), h=Inches(6.1))
    notes(s, "No High class on the open coast. Medium at Bulala Sur/Norte and parts of Linao, Dodan, Maura. Dotted = low confidence, grey = not classified.")

    s = prs.slides.add_slide(blank); title(s, "Rate along the coast, with the error band")
    pic(s, "outputs/maps/Fig_4_3_epr_profile_along_coast.png", Inches(0.5), Inches(1.4), w=Inches(12.3))
    notes(s, f"Grey band = ±{u:.2f} m/yr (one 30 m pixel at both dates, a what-if because no measured error exists). Red dots are the independent set.")

    s = prs.slides.add_slide(blank); title(s, "Table 4.2 — barangay-level results")
    df = f["th"].round(2)
    table(s, df, Inches(0.5), Inches(1.5), Inches(12.3), Inches(3.8), 12)
    add_text(s, Inches(0.5), Inches(5.6), Inches(12.3), Inches(1.5), [
        "Hectares = kilometres of classified shoreline × a 100 m strip (assumption). Kilometres per class are in the extended table.",
        f"Medium: Bulala Sur and Bulala Norte (most robust). Low vs Medium elsewhere is within the ±{u:.2f} m/yr error band for much of the coast."], 14, False, GREY)

    s = prs.slides.add_slide(blank); title(s, "Why the figures differ from the proposal stage")
    add_text(s, Inches(0.6), Inches(1.5), Inches(7.6), Inches(5.5), [
        f"The earlier layers could not be reproduced: {f['audit_hand']} hand-drawn polygons carried all the Medium/High labels; 48 polygons labelled “Medium” had a rate of 0.857 m/yr (< 2).",
        f"The Earth Engine statistics were computed in 1 km circles around points; two study barangays were empty; {f['gee_acc_sea']:.0f} % of “accretion” pixels were open sea.",
        "We rebuilt the analysis from the shoreline lines with a documented pipeline and told the adviser (memo).",
        "Result: High along Bulala (old Fig. 4.2) is replaced by Medium; no High on the open coast."], 17)
    pic(s, "outputs/maps/Fig_legacy_vs_rebuilt.png", Inches(8.5), Inches(1.8), w=Inches(4.6))
    notes(s, "Be transparent. This slide pre-empts the question “why did your map change?”.")

    s = prs.slides.add_slide(blank); title(s, "Survey: ranking of the five structures (Table 4.4)")
    add_text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), [
        "54 experts rated 5 structures on 6 criteria (1–5): Seawall 4.35 (Very High) > Mangrove 4.14 > Revetment 3.99 > Breakwater 3.71 > Riprap 3.63.",
        f"Ranking is robust to criterion weights: Seawall first in {f['ws_first_seawall']:.0f} % of random weightings; the order of the others is not.",
        "Limits: area-wide ranking (not per risk level); 54 respondents, mostly early-career; no SD or test reported yet.",
        "[STUDENTS: add Cronbach’s alpha and the test among structures once the raw data are analysed]"], 19)
    notes(s, "Do not claim a test that was not run. See docs/06_survey_statistics.md.")

    s = prs.slides.add_slide(blank); title(s, "Limitations and next steps")
    add_text(s, Inches(0.6), Inches(1.5), Inches(6.2), Inches(5.3), [
        "Limitations", "• 30 m imagery; dates, tide and measured error not recorded",
        "• Waterline indicator; erosion rate only (no exposure)", "• River mouth not classified", "• No field verification yet", "• Survey not per risk level"], 17)
    add_text(s, Inches(7.0), Inches(1.5), Inches(5.8), Inches(5.3), [
        "Next steps", "• Field check at 20 points (form provided)", "• Seasonal Landsat composites with recorded scenes (script provided)",
        "• Exposure overlay: buildings, roads", "• Per-risk-level questionnaire", "• Field-inspection priority: Bulala Norte, Bulala Sur, Linao"], 17)
    notes(s, "End with the field-inspection priority ranking (outputs/tables/priority_ranking_field_inspection.csv), which is stable across weights.")
    return prs


def run() -> None:  # pragma: no cover
    cfg = load_config()
    out = ensure_dir("outputs/defense")
    prs = build(cfg)
    prs.save(out / "Aparri_GIS_defense_slides.pptx")
    print("deck written to", out)


if __name__ == "__main__":  # pragma: no cover
    run()
