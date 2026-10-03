"""Phase 5: Table 4.2 (barangay-level EPR and risk) and companion tables.

Never copies a number from the legacy CSV (rule: "Never copy numbers from the legacy CSV into Table 4.2").
Everything is computed from the rebuilt transects/segments:

* Mean / Min / Max EPR (m/yr)   over the *valid* transects of the barangay (both end years found, no crossing)
* Low / Medium / High           PRIMARY = kilometres of shoreline in the class (defensible: a shoreline is a line)
                                SECONDARY = hectares of a coastal strip of width ``risk.strip_width_m`` along it
* data-quality columns          transects, valid, flagged, share eroding, mean LRR, confidence mix
"""
from __future__ import annotations

import geopandas as gpd
import numpy as np
import pandas as pd

from .io import REPO_ROOT, ensure_dir, load_config, repo_path
from .risk import CLASSES, length_by_class
from .runreport import RunReport


def barangay_table(tr: pd.DataFrame, seg: gpd.GeoDataFrame, cfg: dict) -> pd.DataFrame:
    order = cfg["table_4_2_order"]
    strip = cfg["risk"].get("strip_width_m")
    km = length_by_class(seg)
    rows = []
    for b in order:
        t = tr[tr.barangay == b]
        v = t[t.valid_epr]
        k = km.loc[b] if b in km.index else pd.Series({c: 0.0 for c in CLASSES + ["No data"]})
        cov = float(k[CLASSES].sum())
        conf = v.confidence.value_counts() if len(v) else pd.Series(dtype=int)
        row = {
            "Barangay": b,
            "Mean EPR (m/yr)": v.EPR_m_yr.mean() if len(v) else np.nan,
            "Min EPR (m/yr)": v.EPR_m_yr.min() if len(v) else np.nan,
            "Max EPR (m/yr)": v.EPR_m_yr.max() if len(v) else np.nan,
            "Low risk (km)": k["Low"], "Medium risk (km)": k["Medium"], "High risk (km)": k["High"],
            "Unclassified (km)": k["No data"], "Classified shoreline (km)": cov,
            "n transects": len(t), "n valid": len(v), "n flagged": int(t.flagged.sum()) if len(t) else 0,
            "% eroding": float((v.EPR_m_yr < 0).mean() * 100) if len(v) else np.nan,
            "Mean LRR (m/yr)": v.lrr.mean() if len(v) else np.nan,
            "confidence high/medium/low (n)": f"{int(conf.get('high', 0))}/{int(conf.get('medium', 0))}/{int(conf.get('low', 0))}",
            "% classes that may flip": float(v.class_may_flip.mean() * 100) if len(v) else np.nan,
        }
        if strip:
            for c in CLASSES:
                row[f"{c} risk (ha, {strip:g} m strip)"] = k[c] * 1000.0 * strip / 1e4
        row["data quality"] = quality_text(t, v, cov)
        rows.append(row)
    df = pd.DataFrame(rows)
    return df


def quality_text(t: pd.DataFrame, v: pd.DataFrame, classified_km: float) -> str:
    if len(t) == 0:
        return "NO DATA: no open-coast transect reaches this barangay"
    bits = [f"{len(v)} of {len(t)} transects valid", f"{classified_km:.2f} km classified"]
    if int(t.flagged.sum()):
        bits.append(f"{int(t.flagged.sum())} flagged")
    return "; ".join(bits)


def thesis_format(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    strip = cfg["risk"].get("strip_width_m")
    cols = ["Barangay", "Mean EPR (m/yr)", "Min EPR (m/yr)", "Max EPR (m/yr)"]
    if strip:
        cols += [f"{c} risk (ha, {strip:g} m strip)" for c in CLASSES]
    else:
        cols += [f"{c} risk (km)" for c in CLASSES]
    out = df[cols].copy()
    out.columns = ["Barangay", "Mean EPR (m/year)", "Minimum EPR (m/year)", "Maximum EPR (m/year)", "Low Risk (ha)", "Medium Risk (ha)", "High Risk (ha)"] if strip \
        else ["Barangay", "Mean EPR (m/year)", "Minimum EPR (m/year)", "Maximum EPR (m/year)", "Low Risk (km)", "Medium Risk (km)", "High Risk (km)"]
    return out


def write_xlsx(path, thesis: pd.DataFrame, ext: pd.DataFrame, notes: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active; ws.title = "Table 4.2"
    thin = Side(style="thin", color="999999")
    def dump(ws, df, title):
        ws.append([title]); ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(list(df.columns))
        for c in ws[3]:
            c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7"); c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for _, r in df.iterrows():
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (round(v, 3) if isinstance(v, float) else v) for v in r.tolist()])
        for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
            for c in row:
                c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if c.column > 1:
                    c.alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 18
        for col in "BCDEFGHIJKLMNOPQRS":
            ws.column_dimensions[col].width = 16
        ws.row_dimensions[3].height = 45
        ws.freeze_panes = "B4"
    dump(ws, thesis, "Table 4.2 Barangay-Level EPR and Coastal Erosion Risk Results, 1990-2025  [DRAFT - FOR STUDENT REVIEW]")
    ws2 = wb.create_sheet("Extended (km, quality)")
    dump(ws2, ext, "Extended version: kilometres of shoreline per class and data-quality columns")
    ws3 = wb.create_sheet("Notes")
    for n in notes:
        ws3.append([n])
    ws3.column_dimensions["A"].width = 140
    wb.save(path)


def write_docx(path, thesis: pd.DataFrame, notes: list[str]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    p = doc.add_paragraph("[DRAFT – FOR STUDENT REVIEW]"); p.runs[0].bold = True
    cap = doc.add_paragraph("Table 4.2 Barangay-Level EPR and Coastal Erosion Risk Results, 1990–2025"); cap.runs[0].bold = True
    t = doc.add_table(rows=1, cols=len(thesis.columns)); t.style = "Table Grid"
    for i, c in enumerate(thesis.columns):
        cell = t.rows[0].cells[i]; cell.text = ""
        r = cell.paragraphs[0].add_run(str(c)); r.bold = True; r.font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _, row in thesis.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row.tolist()):
            txt = "–" if (isinstance(v, float) and np.isnan(v)) else (f"{v:.2f}" if isinstance(v, float) else str(v))
            cells[i].text = txt
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            for r in cells[i].paragraphs[0].runs:
                r.font.size = Pt(9)
    for n in notes[:4]:
        q = doc.add_paragraph(n); q.runs[0].italic = True; q.runs[0].font.size = Pt(9)
    doc.save(path)


def run() -> None:  # pragma: no cover
    cfg = load_config()
    rep = RunReport("tables", cfg)
    name = cfg["shoreline_set"]
    seg = gpd.read_file(repo_path(f"data/processed/{name}/risk_segments.gpkg"))
    tr = gpd.read_file(repo_path(f"data/processed/{name}/risk_transects.gpkg"))
    rep.add_input(f"data/processed/{name}/risk_segments.gpkg"); rep.add_input(f"data/processed/{name}/risk_transects.gpkg")
    df = barangay_table(tr, seg, cfg)
    th = thesis_format(df, cfg)
    strip = cfg["risk"]["strip_width_m"]
    notes = [
        f"Shoreline set: {name}. EPR = (position 2025 − position 1990) ÷ 35 years; negative = erosion, positive = accretion.",
        f"Classes (Table 3.1): Low < 2, Medium 2–5 (inclusive), High > 5 m/year of erosion; accreting transects are counted as Low and tagged 'accreting'.",
        f"Hectares = kilometres of classified shoreline × a {strip:g} m coastal strip (config risk.strip_width_m, decision D-08). The primary, assumption-free measure is kilometres of shoreline.",
        "Only the open, sea-facing coast is classified (the Cagayan River banks are excluded, decision D-06). Bulala Sur and Paddaya are covered because the analysis is not limited to the old Earth Engine box.",
        "Mean/Min/Max use valid transects only (both end years found, no crossing). Acquisition dates and positional error were not supplied: EPR uses 35 calendar years and no measured uncertainty (a ±30 m what-if is used only for the confidence column).",
        "No number in this table comes from the legacy CSV/shapefile.",
    ]
    out = ensure_dir("outputs/tables")
    th.round(3).to_csv(out / "Table_4_2_barangay_results.csv", index=False)
    df.round(3).to_csv(out / "Table_4_2_extended.csv", index=False)
    write_xlsx(out / "Table_4_2_barangay_results.xlsx", th.round(3), df.round(3), notes)
    write_docx(out / "Table_4_2_barangay_results.docx", th, notes)
    # one table per set, for the comparison in Phase 5/6
    from .transects import available_sets
    allsets = []
    for s in available_sets():
        sg = gpd.read_file(repo_path(f"data/processed/{s}/risk_segments.gpkg")); tt = gpd.read_file(repo_path(f"data/processed/{s}/risk_transects.gpkg"))
        d = barangay_table(tt, sg, cfg); d.insert(0, "shoreline_set", s); allsets.append(d)
    pd.concat(allsets).round(3).to_csv(out / "Table_4_2_by_shoreline_set.csv", index=False)
    for f in ("Table_4_2_barangay_results.csv", "Table_4_2_barangay_results.xlsx", "Table_4_2_barangay_results.docx", "Table_4_2_extended.csv", "Table_4_2_by_shoreline_set.csv"):
        rep.add_output(out / f)
    rep.write()
    print(th.round(2).to_string(index=False))
    print(df[["Barangay", "Classified shoreline (km)", "n valid", "n flagged", "% classes that may flip", "confidence high/medium/low (n)"]].round(1).to_string(index=False))


if __name__ == "__main__":  # pragma: no cover
    run()
